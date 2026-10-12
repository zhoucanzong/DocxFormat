from pathlib import Path
import subprocess
import sys

import pytest
from docx import Document

from docxformat.markdown import generate_baseline, parse_markdown
from docxformat.package import FormatError, Package, NS, qn
from docxformat.verification import verify_package


def test_md_keeps_deep_headings_lists_tables_links_quotes_and_code(tmp_path):
    src = tmp_path / "body.md"
    src.write_text("# 文档\n\n#### 四级标题\n\n3. 第一项\n   - 子项\n4. 第二项\n\n> 引用内容\n\n```python\nx = 1\n\n    y = 2\n```\n\n| 名称 | 数值 |\n| --- | ---: |\n| **样本** | 2.500 |\n\n[外部链接](https://example.com)\n", encoding="utf-8")
    out = tmp_path / "body.docx"
    report = generate_baseline(src, out)
    assert report["counts"]["tables"] == 1
    doc = Document(out)
    assert next(p for p in doc.paragraphs if p.text == "四级标题").style.name == "Heading 4"
    assert next(p for p in doc.paragraphs if p.text == "引用内容").style.name == "Quote"
    code = next(p for p in doc.paragraphs if p.style.name == "Code Block")
    assert code.text == "x = 1\n\n    y = 2\n"
    assert doc.tables[0].cell(1, 0).paragraphs[0].runs[0].bold
    assert doc.tables[0].cell(1, 1).text == "2.500"
    pkg = Package(out)
    nums = pkg.root("word/numbering.xml")
    assert nums.xpath(".//w:lvl/w:start[@w:val='3']", namespaces=NS)
    assert pkg.root("word/document.xml").xpath(".//w:ilvl[@w:val='1']", namespaces=NS)
    assert pkg.root("word/document.xml").xpath(".//w:hyperlink", namespaces=NS)
    assert not verify_package(out)["errors"]


@pytest.mark.parametrize("text", ["<div>raw HTML</div>", "正文 <span>内嵌HTML</span>"])
def test_unimplemented_md_extensions_are_not_silently_lost(tmp_path, text):
    src, out = tmp_path / "input.md", tmp_path / "out.docx"
    src.write_text(text)
    assert parse_markdown(src)["unsupported"]
    with pytest.raises(FormatError, match="unsupported"):
        generate_baseline(src, out)
    assert not out.exists()


def test_missing_image_does_not_publish_partial_output(tmp_path):
    src, out = tmp_path / "input.md", tmp_path / "out.docx"
    src.write_text("# 标题\n\n![说明](missing.png)")
    with pytest.raises(FormatError, match="Missing Markdown image"):
        generate_baseline(src, out)
    assert not out.exists()


def test_portable_entrypoint_works_outside_repository(tmp_path):
    script = Path(__file__).resolve().parents[1] / "skills/docx-format/scripts/docx_format.py"
    run = subprocess.run([sys.executable, str(script), "--help"], cwd=tmp_path, text=True, capture_output=True)
    assert run.returncode == 0, run.stderr
    assert "approve-plan" in run.stdout


def test_notes_controls_internal_links_and_editable_math(tmp_path):
    src, out = tmp_path / "input.md", tmp_path / "out.docx"
    src.write_text("# 第一章\n\n[跳到章节](#第一章)，正文[^n]含公式 $x^2+1$。\n\n- [x] 已检查\n- [ ] 待检查\n\n$$\n\\frac{a+b}{c}\n$$\n\n[^n]: 脚注含[链接](https://example.com)。\n")
    result = generate_baseline(src, out)
    assert result["counts"]["footnotes"] == 1
    assert result["counts"]["task_boxes"] == 2
    assert result["counts"]["equations"] == 2
    pkg = Package(out)
    root = pkg.root("word/document.xml")
    assert len(root.xpath(".//m:oMath", namespaces=NS)) == 2
    assert len(root.xpath(".//m:f", namespaces=NS)) == 1
    assert len(root.xpath(".//m:sSup", namespaces=NS)) == 1
    names = root.xpath(".//w:bookmarkStart/@w:name", namespaces=NS)
    assert root.xpath(".//w:hyperlink/@w:anchor", namespaces=NS) == names
    assert root.xpath(".//w:footnoteReference/@w:id", namespaces=NS) == ["1"]
    notes = pkg.root("word/footnotes.xml")
    assert notes.xpath("./w:footnote[@w:id='1']//w:footnoteRef", namespaces=NS)
    assert notes.xpath(".//w:hyperlink", namespaces=NS)
    assert not verify_package(out)["errors"]


@pytest.mark.parametrize("text", ["[坏链接](#不存在)", r"数学 $\unknowncommand{x}$"])
def test_unresolved_anchors_and_math_commands_fail_before_publication(tmp_path, text):
    src, out = tmp_path / "input.md", tmp_path / "out.docx"
    src.write_text(text)
    with pytest.raises(FormatError):
        generate_baseline(src, out)
    assert not out.exists()
