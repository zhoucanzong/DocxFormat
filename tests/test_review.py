from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as dq
from docx.shared import Pt
from PIL import Image

from docxformat.engine import apply_plan, approve_plan
from docxformat.inspection import inspect_document
from docxformat.package import FormatError, NS, Package, digest, protected_signature, qn
from docxformat.planning import build_plan, extract_template
from docxformat.revisions import resolve_review, revision_inventory, tree_digest


def fixture(path):
    doc = Document()
    doc.add_heading("测试论文", 0)
    p = doc.add_paragraph()
    p.add_run("保留普通文字，").bold = False
    p.add_run("这是重点。").bold = True
    doc.add_heading("第一章 示例", 1)
    p = doc.add_paragraph("公式区域：")
    math = OxmlElement("m:oMath")
    mr = OxmlElement("m:r")
    mt = OxmlElement("m:t")
    mt.text = "x=1"
    mr.append(mt)
    math.append(mr)
    p._p.append(math)
    image = BytesIO()
    Image.new("RGB", (24, 24), "red").save(image, format="PNG")
    image.seek(0)
    doc.add_picture(image)
    table = doc.add_table(rows=4, cols=2)
    for i, row in enumerate(table.rows):
        row.cells[0].text = ["项目", "单位", "样本A", "样本B"][i]
        row.cells[1].text = ["结果", "mm", "2.500", "3.200"][i]
    table.cell(0, 0).merge(table.cell(0, 1))
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("第二节正文。")
    p = doc.sections[0].footer.paragraphs[0]
    p.add_run("保留页脚说明 ")
    for tag, value in [("fldChar", "begin"), ("instrText", " PAGE "), ("fldChar", "separate"), ("t", "1"), ("fldChar", "end")]:
        run = p.add_run()
        el = OxmlElement("w:" + tag)
        if tag == "fldChar":
            el.set(dq("w:fldCharType"), value)
        else:
            el.text = value
        run._r.append(el)
    # An existing editorial insertion must remain pending and byte-equivalent.
    p = doc.add_paragraph()
    ins = OxmlElement("w:ins")
    for k, v in {"id": "71", "author": "Original reviewer", "date": "2025-01-01T00:00:00Z"}.items():
        ins.set(dq("w:" + k), v)
    r, t = OxmlElement("w:r"), OxmlElement("w:t")
    t.text = "原有待审内容"
    r.append(t)
    ins.append(r)
    p._p.append(ins)
    doc.save(path)


def profile():
    return {"schema_version": 1, "name": "Test rules", "roles": {
        "body": {"paragraph": {"line": {"mode": "multiple", "value": 1.5}, "first_line_chars": 2},
                 "run": {"font_east_asia": "宋体", "font_latin": "Times New Roman", "size_pt": 12, "bold": False}},
        "heading1": {"paragraph": {"keep_with_next": True}, "run": {"size_pt": 16}},
    }, "table_rules": [{"index": 0, "header_rows": 2}],
        "section_rules": [{"index": 0, "properties": {"top_mm": 30}}]}


def review(tmp_path):
    source = tmp_path / "source.docx"
    fixture(source)
    plan = build_plan(source, profile())
    approval = approve_plan(plan, confirmation="User approved this exact test plan")
    output = tmp_path / "review.docx"
    receipt = apply_plan(plan, approval, output, tmp_path / "receipt.json")
    return source, plan, output, receipt


def test_tracked_accept_reject_preserve_content_objects_and_history(tmp_path):
    source, plan, output, receipt = review(tmp_path)
    before = Package(source)
    inspected = Package(output)
    assert receipt["changes"]
    assert any(c["tag"] == "tcPrChange" for c in receipt["changes"])
    assert any(c["tag"] == "sectPrChange" for c in receipt["changes"])
    assert protected_signature(before) == protected_signature(inspected)
    assert before.source_hash == plan["source_sha256"]
    for name in before.parts:
        if name not in {c["part"] for c in receipt["changes"]} | {"word/settings.xml"}:
            assert before.parts[name] == inspected.parts[name]
    assert inspected.root("word/settings.xml").find(qn("w:trackRevisions")) is not None
    accepted, rejected = tmp_path / "accepted.docx", tmp_path / "rejected.docx"
    resolve_review(output, receipt, accepted, action="accept", confirmation="Accept this round")
    resolve_review(output, receipt, rejected, action="reject", confirmation="Reject this round")
    for file in [accepted, rejected]:
        pkg = Package(file)
        assert protected_signature(pkg) == protected_signature(before)
        assert revision_inventory(pkg) == revision_inventory(before)
        assert pkg.parts["word/settings.xml"] == before.parts["word/settings.xml"]
    # Reject returns old properties and preserves the section inside pPr.
    assert tree_digest(Package(rejected).root("word/document.xml")) == tree_digest(before.root("word/document.xml"))
    # Mixed inline emphasis survives accepting body normalization.
    accepted_doc = Document(accepted)
    p = next(p for p in accepted_doc.paragraphs if "保留普通文字" in p.text)
    assert [r.bold for r in p.runs] == [False, True]
    assert all(r.font.size.pt == 12 for r in p.runs)
    assert accepted_doc.sections[0].top_margin.mm == pytest.approx(30, abs=0.05)
    table = accepted_doc.tables[0]
    def bottom(row):
        return table.rows[row].cells[0]._tc.find("w:tcPr/w:tcBorders/w:bottom", NS)
    assert bottom(0).get(qn("w:val")) == "nil"
    assert bottom(1).get(qn("w:sz")) == "4"
    assert bottom(3).get(qn("w:sz")) == "12"


def test_gate_rejects_missing_confirmation_changed_plan_or_changed_source(tmp_path):
    src = tmp_path / "source.docx"
    fixture(src)
    plan = build_plan(src, profile())
    with pytest.raises(FormatError, match="confirmation"):
        approve_plan(plan, confirmation="")
    approval = approve_plan(plan, confirmation="Approved")
    tampered = deepcopy(plan)
    tampered["operations"][0]["properties"] = {"size_pt": 22}
    with pytest.raises(FormatError, match="digest"):
        apply_plan(tampered, approval, tmp_path / "out.docx", tmp_path / "receipt.json")
    doc = Document(src)
    doc.add_paragraph("用户改了原稿")
    doc.save(src)
    with pytest.raises(FormatError, match="Source changed"):
        apply_plan(plan, approval, tmp_path / "out.docx", tmp_path / "receipt.json")


def test_partial_accept_never_accepts_old_revisions(tmp_path):
    source, plan, output, receipt = review(tmp_path)
    partial = tmp_path / "partial.docx"
    first = receipt["changes"][0]["id"]
    result = resolve_review(output, receipt, partial, action="accept", confirmation="Accept one", selected_ids=[first])
    assert len(result["changes"]) == len(receipt["changes"]) - 1
    assert result["state"] == "partially-approved"
    final = tmp_path / "after.docx"
    resolve_review(partial, result, final, action="accept", confirmation="Accept remaining")
    assert revision_inventory(Package(final)) == revision_inventory(Package(source))


def test_review_cannot_be_resolved_silently_or_after_manual_edit(tmp_path):
    _, _, output, receipt = review(tmp_path)
    with pytest.raises(FormatError, match="confirmation"):
        resolve_review(output, receipt, tmp_path / "a.docx", action="accept", confirmation="")
    doc = Document(output)
    doc.add_paragraph("审阅后增加内容")
    doc.save(output)
    with pytest.raises(FormatError, match="changed after"):
        resolve_review(output, receipt, tmp_path / "a.docx", action="accept", confirmation="Accept")


def test_template_extraction_exposes_conflicts_and_protects_specials(tmp_path):
    src = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("正文一").runs[0].font.size = Pt(12)
    doc.add_paragraph("正文二").runs[0].font.size = Pt(14)
    doc.save(src)
    result = extract_template(src)
    assert any(q["id"] == "body.run:size_pt" for q in result["profile"]["questions"])
    assert "size_pt" not in result["profile"]["roles"]["body"]["run"]


def test_untouched_noop_does_not_create_revisions(tmp_path):
    source, _, output, receipt = review(tmp_path)
    accepted = tmp_path / "accepted.docx"
    resolve_review(output, receipt, accepted, action="accept", confirmation="Accept")
    plan = build_plan(accepted, profile())
    approval = approve_plan(plan, confirmation="Apply same rules")
    next_receipt = apply_plan(plan, approval, tmp_path / "review2.docx", tmp_path / "receipt2.json")
    assert not next_receipt["changes"]


def test_source_version_is_checked_again_before_publishing(tmp_path):
    source = tmp_path / "source.docx"
    fixture(source)
    pkg = Package(source)
    doc = Document(source)
    doc.add_paragraph("Concurrent author edit")
    doc.save(source)
    with pytest.raises(FormatError, match="during execution"):
        pkg.save(tmp_path / "stale.docx")
    assert not (tmp_path / "stale.docx").exists()


def test_section_numbering_columns_and_first_page_are_reversible(tmp_path):
    source = tmp_path / "source.docx"
    fixture(source)
    rules = {"schema_version": 1, "section_rules": [{"index": 1, "properties": {
        "orientation": "landscape", "width_mm": 297, "height_mm": 210,
        "first_page_different": True, "page_number_start": 1,
        "page_number_format": "upperRoman", "column_count": 2, "column_space_mm": 8}}]}
    plan = build_plan(source, rules)
    review_path = tmp_path / "review.docx"
    receipt = apply_plan(plan, approve_plan(plan, confirmation="Approve specified section only"), review_path, tmp_path / "receipt.json")
    from docxformat.verification import verify_package
    assert verify_package(review_path, original=source, receipt=receipt, plan=plan)["rule_check"]["status"] == "passed"
    output = tmp_path / "reject.docx"
    resolve_review(review_path, receipt, output, action="reject", confirmation="Reject section change")
    assert tree_digest(Package(source).root("word/document.xml")) == tree_digest(Package(output).root("word/document.xml"))
