from io import BytesIO
from copy import deepcopy

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import RGBColor, Pt
from PIL import Image
from lxml import etree

from docxformat.components import approve_component, catalog, component_plan, preview_component
from docxformat.package import FormatError, Package, NS, digest, element, qn
from docxformat.revisions import tree_digest
from docxformat.verification import verify_package


def base(path):
    doc = Document()
    doc.add_heading("第一章", 1)
    doc.add_paragraph("原文保持不变。")
    doc.save(path)


def template(path):
    doc = Document()
    style = doc.styles.add_style("AgencyMark", WD_STYLE_TYPE.PARAGRAPH)
    style.font.color.rgb = RGBColor(200, 0, 0)
    style.font.size = Pt(28)
    p = doc.add_paragraph("示例机关文件", style=style)
    rpr = p.runs[0]._r.get_or_add_rPr()
    scale = etree.SubElement(rpr, qn("w:w"))
    scale.set(qn("w:val"), "80")
    p = doc.add_paragraph()
    stream = BytesIO()
    Image.new("RGB", (120, 3), (200, 0, 0)).save(stream, "PNG")
    stream.seek(0)
    p.add_run().add_picture(stream)
    doc.add_paragraph("不要复制的模板正文")
    doc.save(path)


def test_letterhead_import_preserves_native_scale_media_and_unselected_body(tmp_path):
    source, tpl = tmp_path / "source.docx", tmp_path / "template.docx"
    base(source)
    template(tpl)
    src_hash, tpl_hash = digest(source.read_bytes()), digest(tpl.read_bytes())
    dst_blocks = catalog(source)["blocks"]
    blocks = catalog(tpl)["blocks"][:2]
    plan = component_plan(source, {"type": "letterhead", "template": str(tpl),
                                  "block_ids": [b["id"] for b in blocks], "before_id": dst_blocks[0]["id"]})
    preview = tmp_path / "preview.docx"
    receipt = preview_component(plan, preview, tmp_path / "preview.json")
    assert receipt["state"] == "awaiting-component-approval" and not receipt["native_revision"]
    pkg = Package(preview)
    body = pkg.root("word/document.xml").find(qn("w:body"))
    oldbody = Package(source).root("word/document.xml").find(qn("w:body"))
    assert [tree_digest(n) for n in body[2:]] == [tree_digest(n) for n in oldbody]
    assert pkg.root("word/document.xml").xpath(".//w:rPr/w:w[@w:val='80']", namespaces=NS)
    assert "不要复制的模板正文" not in "".join(pkg.root("word/document.xml").xpath(".//w:t/text()", namespaces=NS))
    media = [digest(v) for k, v in pkg.parts.items() if k.startswith("word/media/")]
    assert media == [digest(v) for k, v in Package(tpl).parts.items() if k.startswith("word/media/")]
    assert not verify_package(preview)["errors"]
    assert digest(source.read_bytes()) == src_hash and digest(tpl.read_bytes()) == tpl_hash
    with pytest.raises(FormatError, match="approval"):
        approve_component(receipt, tmp_path / "final.docx", confirmation="")
    result = approve_component(receipt, tmp_path / "final.docx", confirmation="User approved this component preview")
    assert result["state"] == "component-approved"
    assert (tmp_path / "final.docx").read_bytes() == preview.read_bytes()


def test_toc_creates_a_real_field_without_invented_page_numbers(tmp_path):
    src = tmp_path / "source.docx"
    base(src)
    first = catalog(src)["blocks"][0]["id"]
    plan = component_plan(src, {"type": "toc", "levels": [1, 4], "before_id": first})
    out = tmp_path / "preview.docx"
    receipt = preview_component(plan, out, tmp_path / "preview.json")
    root = Package(out).root("word/document.xml")
    assert ' TOC \\o "1-4" \\h \\z \\u ' in root.xpath(".//w:instrText/text()", namespaces=NS)
    assert receipt["fields"] == "requires-engine-update"
    assert "原文保持不变。" in root.xpath(".//w:t/text()", namespaces=NS)
    with pytest.raises(FormatError, match="Existing TOC"):
        component_plan(out, {"type": "toc"})


def test_component_approval_rejects_modified_preview(tmp_path):
    src = tmp_path / "source.docx"
    base(src)
    plan = component_plan(src, {"type": "toc"})
    out = tmp_path / "preview.docx"
    receipt = preview_component(plan, out, tmp_path / "preview.json")
    doc = Document(out)
    doc.add_paragraph("Additional manual change")
    doc.save(out)
    with pytest.raises(FormatError, match="changed"):
        approve_component(receipt, tmp_path / "final.docx", confirmation="Approve older preview")


def test_component_never_deletes_section_or_revision_blocks(tmp_path):
    from test_review import fixture
    src = tmp_path / "source.docx"
    fixture(src)
    observed = catalog(src)
    blocked = [b["id"] for b in observed["blocks"] if b["has_section"] or b["existing_revisions"]]
    assert len(blocked) >= 2
    for ident in blocked:
        with pytest.raises(FormatError, match="historical revisions or section"):
            component_plan(src, {"type": "toc", "replace_ids": [ident]})


def test_vml_wordart_type_reference_is_migrated_across_selected_blocks(tmp_path):
    source, tpl = tmp_path / "source.docx", tmp_path / "effects.docx"
    base(source)
    doc = Document()
    pict = element("w:pict")
    shapetype = element("v:shapetype")
    shapetype.set("id", "_x0000_t136")
    shapetype.set("coordsize", "21600,21600")
    pict.append(shapetype)
    doc.add_paragraph().add_run()._r.append(pict)
    pict = element("w:pict")
    shape = element("v:shape")
    shape.set("id", "_x0000_s1025")
    shape.set("type", "#_x0000_t136")
    shape.set("style", "width:300pt;height:40pt")
    shape.set("fillcolor", "#cc0000")
    textpath = element("v:textpath")
    textpath.set("string", "特殊红头")
    textpath.set("style", "font-family:宋体;font-size:28pt")
    shape.append(textpath); pict.append(shape)
    doc.add_paragraph().add_run()._r.append(pict)
    doc.save(tpl)
    selected = [b["id"] for b in catalog(tpl)["blocks"]]
    first = catalog(source)["blocks"][0]["id"]
    plan = component_plan(source, {"type": "letterhead", "template": str(tpl), "block_ids": selected, "before_id": first})
    output = tmp_path / "preview.docx"
    preview_component(plan, output, tmp_path / "receipt.json")
    root = Package(output).root("word/document.xml")
    migrated_type = root.xpath(".//v:shapetype/@id", namespaces=NS)[0]
    assert root.xpath(".//v:shape/@type", namespaces=NS) == ["#" + migrated_type]
    assert root.xpath(".//v:textpath/@string", namespaces=NS) == ["特殊红头"]
    assert root.xpath(".//v:shape/@style", namespaces=NS) == ["width:300pt;height:40pt"]


def test_imported_theme_style_is_isolated_from_destination_theme(tmp_path):
    source, tpl = tmp_path / "source.docx", tmp_path / "themed.docx"
    base(source)
    doc = Document()
    doc.add_heading("模板标题", 1)
    doc.save(tpl)
    selected = catalog(tpl)["blocks"][0]["id"]
    plan = component_plan(source, {"type": "letterhead", "template": str(tpl), "block_ids": [selected]})
    output = tmp_path / "preview.docx"
    preview_component(plan, output, tmp_path / "receipt.json")
    root = Package(output).root("word/styles.xml")
    imported = root.xpath("./w:style[starts-with(@w:styleId, 'DF')]", namespaces=NS)
    assert imported
    assert not any(s.xpath(".//*[@w:themeColor]", namespaces=NS) for s in imported)
    # Original Heading1 is not overwritten by the imported template style.
    before = Package(source).root("word/styles.xml").xpath("./w:style[@w:styleId='Heading1']", namespaces=NS)[0]
    after = root.xpath("./w:style[@w:styleId='Heading1']", namespaces=NS)[0]
    assert tree_digest(before) == tree_digest(after)
