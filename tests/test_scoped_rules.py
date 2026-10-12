from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as dq
from docxformat.engine import approve_plan, apply_plan
from docxformat.inspection import inspect_document
from docxformat.package import FormatError, Package, NS, element, protected_signature, qn
from docxformat.planning import build_plan, validate_profile
from docxformat.revisions import resolve_review, tree_digest
from docxformat.verification import verify_package


def test_explicit_formula_and_textbox_typography_preserves_objects_and_rejects(tmp_path):
    src = tmp_path / "source.docx"
    doc = Document()
    p = doc.add_paragraph("原生公式：")
    math, mr, mt = element("m:oMath"), element("m:r"), element("m:t")
    mt.text = "x+1"
    mr.append(mt); math.append(mr); p._p.append(math)
    p = doc.add_paragraph()
    pict, shape, box, content = element("w:pict"), element("v:shape"), element("v:textbox"), element("w:txbxContent")
    shape.set("id", "redhead-shape")
    shape.set("style", "width:240pt;height:50pt")
    textp, run, text = element("w:p"), element("w:r"), element("w:t")
    text.text = "原生文本框"
    run.append(text); textp.append(run); content.append(textp); box.append(content); shape.append(box); pict.append(shape)
    p.add_run()._r.append(pict)
    doc.save(src)
    obs = inspect_document(src)
    overrides = {}
    for record in obs["paragraphs"]:
        if "content-control-or-textbox" in record["protection"]:
            overrides[record["id"]] = {"run": {"size_pt": 20}, "reason": "Confirmed textbox typography only"}
        elif "special-object" in record["protection"] and "公式" in record["text"]:
            overrides[record["id"]] = {"paragraph": {"alignment": "center"}, "reason": "Confirmed formula alignment only"}
    assert len(overrides) == 2
    plan = build_plan(src, {"schema_version": 1, "overrides": overrides})
    out = tmp_path / "review.docx"
    receipt = apply_plan(plan, approve_plan(plan, confirmation="User approved two scoped changes"), out, tmp_path / "receipt.json")
    assert verify_package(out, original=src, receipt=receipt, plan=plan)["rule_check"]["status"] == "passed"
    rejected = tmp_path / "rejected.docx"
    resolve_review(out, receipt, rejected, action="reject", confirmation="Reject scoped changes")
    assert tree_digest(Package(src).root("word/document.xml")) == tree_digest(Package(rejected).root("word/document.xml"))


@pytest.mark.parametrize("profile", [
    {"schema_version": 1, "section_rules": "all"},
    {"schema_version": 1, "section_rules": [{"index": True, "properties": {}}]},
    {"schema_version": 1, "table_rules": [None]},
    {"schema_version": 1, "overrides": {"x": {"run": {"size_pt": 12}}}},
])
def test_malformed_profiles_fail_as_actionable_errors(profile):
    with pytest.raises(FormatError):
        validate_profile(profile)
