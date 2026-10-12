"""Markdown AST to editable Word structure, without content rewriting."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from urllib.parse import unquote, urlparse

import mistune
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.packuri import PackURI
from docx.parts.story import StoryPart
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .package import FormatError, NS, atomic_bytes, digest, element, parse_xml, qn

PLUGINS = ["table", "footnotes", "math", "strikethrough", "url", "task_lists"]
SUPPORTED = {"blank_line", "heading", "paragraph", "block_text", "block_code", "block_quote", "list", "list_item",
             "table", "table_head", "table_body", "table_row", "table_cell", "text", "strong", "emphasis",
             "strikethrough", "codespan", "link", "image", "softbreak", "linebreak", "thematic_break",
             "footnotes", "footnote_item", "footnote_ref", "task_list_item", "inline_math", "block_math"}


def native_math(latex):
    try:
        from latex2mathml.converter import convert
        import mathml2omml
    except ImportError as exc:
        raise FormatError("Install the skill requirements for editable equation generation") from exc
    if len(latex) > 50000:
        raise FormatError("Equation exceeds conversion limits")
    try:
        mathml = convert(latex)
        root = parse_xml(mathml.encode())
        if re.search(r"\\[A-Za-z]+", "".join(root.itertext())):
            raise FormatError("Unrecognized LaTeX command; review the equation rather than silently typesetting it as text")
        omml = mathml2omml.convert(mathml)
        wrapper = parse_xml((f'<root xmlns:m="{NS["m"]}" xmlns:w="{NS["w"]}">' + omml + "</root>").encode())
        if len(wrapper) != 1 or wrapper[0].tag != qn("m:oMath"):
            raise FormatError("Equation converter did not return a native math object")
        return wrapper[0]
    except FormatError:
        raise
    except Exception as exc:
        raise FormatError(f"Equation requires manual review: {exc}") from exc


class NoteContainer:
    def __init__(self, root, part, doc):
        self.root, self.part, self.doc = root, part, doc

    def add_paragraph(self, text="", style=None):
        node = OxmlElement("w:p")
        self.root.append(node)
        p = Paragraph(node, self)
        if style:
            p.style = style
        if text:
            p.add_run(text)
        return p

    def add_table(self, rows, cols):
        table = self.doc.add_table(rows=rows, cols=cols)
        table._tbl.getparent().remove(table._tbl)
        self.root.append(table._tbl)
        return Table(table._tbl, self)


def parse_markdown(path):
    source = Path(path).read_text(encoding="utf-8-sig")
    if source.startswith("---\n"):
        raise FormatError("Markdown front matter must first be separated into confirmed metadata and body")
    ast = mistune.create_markdown(renderer=None, plugins=PLUGINS)(source)
    unsupported = []
    def visit(nodes):
        for node in nodes:
            if node["type"] not in SUPPORTED:
                unsupported.append({"type": node["type"], "sample": node.get("raw", "")[:100]})
            visit(node.get("children", []))
    visit(ast)
    return {"source": str(Path(path).resolve()), "sha256": digest(Path(path).read_bytes()), "ast": ast, "unsupported": unsupported}


class MarkdownWriter:
    def __init__(self, source, *, title_from_h1=False):
        self.source = Path(source).resolve()
        self.doc = Document()
        self.title_from_h1 = title_from_h1
        self.title_used = False
        self.counts = {"headings": 0, "lists": 0, "tables": 0, "links": 0, "images": 0, "code_blocks": 0, "quotes": 0, "footnotes": 0, "task_boxes": 0, "equations": 0}
        self.bookmarks = {}
        self.note_part = None
        self.bookmark_number = 1
        self.doc.sections[0].page_width = Inches(8.2677)
        self.doc.sections[0].page_height = Inches(11.6929)
        self.doc.styles["Normal"].paragraph_format.space_after = Pt(6)
        # This is an unstyled semantic baseline; target typography is applied
        # as tracked changes by the same planner as existing DOCX inputs.
        for name in ["Code Block", "Code Inline"]:
            if name not in self.doc.styles:
                style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH if name == "Code Block" else WD_STYLE_TYPE.CHARACTER)
                style.font.name = "Courier New"
                style.font.size = Pt(10)
        self.doc.styles["Code Block"].paragraph_format.first_line_indent = Pt(0)
        for name in ["Title", "Subtitle", *[f"Heading {i}" for i in range(1, 10)]]:
            style = self.doc.styles[name]
            style.font.color.rgb = RGBColor(0, 0, 0)
            if style.element.pPr is not None:
                for border in list(style.element.pPr.findall(qn("w:pBdr"))):
                    style.element.pPr.remove(border)

    def prepare(self, nodes):
        seen = {}
        def visit(items):
            for node in items:
                if node["type"] == "heading":
                    raw = unicodedata.normalize("NFC", self.plain(node.get("children", []))).lower()
                    slug = re.sub(r"\s", "-", re.sub(r"[^\w\s-]", "", raw).strip())
                    count = seen.get(slug, 0)
                    seen[slug] = count + 1
                    slug = slug + ("-" + str(count) if count else "")
                    name = "DF_" + digest(slug.encode())[:24]
                    self.bookmarks[slug] = name
                    node["_bookmark"] = name
                visit(node.get("children", []))
        visit(nodes)

    def ensure_notes(self):
        if self.note_part is not None:
            return self.note_part
        if "Footnote Text" not in self.doc.styles:
            style = self.doc.styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
            style.font.size = Pt(10)
            style.paragraph_format.space_after = Pt(0)
        if "Footnote Reference" not in self.doc.styles:
            style = self.doc.styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
            style.font.superscript = True
        root = OxmlElement("w:footnotes")
        for ident, kind in [(-1, "separator"), (0, "continuationSeparator")]:
            note = element("w:footnote", id=ident, type=kind)
            p, r = element("w:p"), element("w:r")
            r.append(element("w:" + kind))
            p.append(r)
            note.append(p)
            root.append(note)
        self.note_part = StoryPart(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, root, self.doc.part.package)
        self.doc.part.relate_to(self.note_part, RT.FOOTNOTES)
        return self.note_part

    def numbering(self, ordered, start):
        root = self.doc.part.numbering_part.element
        abstract_id = max([int(n.get(qn("w:abstractNumId"))) for n in root.findall(qn("w:abstractNum"))], default=-1) + 1
        num_id = max([int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))], default=0) + 1
        abstract = element("w:abstractNum", abstractNumId=abstract_id)
        abstract.append(element("w:multiLevelType", val="multilevel"))
        for depth in range(9):
            lvl = element("w:lvl", ilvl=depth)
            lvl.append(element("w:start", val=start))
            lvl.append(element("w:numFmt", val="decimal" if ordered else "bullet"))
            lvl.append(element("w:lvlText", val=f"%{depth + 1}." if ordered else "•"))
            lvl.append(element("w:suff", val="tab"))
            ppr = element("w:pPr")
            tabs = element("w:tabs")
            tabs.append(element("w:tab", val="num", pos=720 * (depth + 1)))
            ppr.append(tabs)
            ppr.append(element("w:ind", left=720 * (depth + 1), hanging=360))
            lvl.append(ppr)
            abstract.append(lvl)
        # All abstract definitions precede numbering instances.
        first_num = root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            root.append(abstract)
        num = element("w:num", numId=num_id)
        num.append(element("w:abstractNumId", val=abstract_id))
        root.append(num)
        self.counts["lists"] += 1
        return num_id

    def inline(self, p, nodes, attrs=None):
        attrs = attrs or {}
        for node in nodes:
            typ = node["type"]
            if typ in {"strong", "emphasis", "strikethrough"}:
                key = {"strong": "bold", "emphasis": "italic", "strikethrough": "strike"}[typ]
                self.inline(p, node["children"], {**attrs, key: True})
            elif typ in {"text", "codespan", "softbreak", "linebreak"}:
                text = " " if typ == "softbreak" else "\n" if typ == "linebreak" else node.get("raw", "")
                run = p.add_run(text)
                for key, value in attrs.items():
                    setattr(run.font, key, value)
                if typ == "codespan":
                    run.style = self.doc.styles["Code Inline"]
            elif typ == "link":
                url = node["attrs"]["url"]
                scheme = urlparse(url).scheme.lower()
                if scheme and scheme not in {"http", "https", "mailto", "file"}:
                    raise FormatError(f"Unsupported link scheme: {scheme}")
                if url.startswith("#"):
                    anchor = unquote(url[1:])
                    if anchor not in self.bookmarks:
                        raise FormatError(f"Markdown link has no matching heading: {anchor}")
                    link = element("w:hyperlink", anchor=self.bookmarks[anchor])
                else:
                    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
                    link = element("w:hyperlink")
                    link.set(qn("r:id"), rid)
                temp = Paragraph(OxmlElement("w:p"), p._parent)
                self.inline(temp, node.get("children", []), attrs)
                for run in list(temp._p):
                    if run.tag != qn("w:pPr"):
                        link.append(run)
                p._p.append(link)
                self.counts["links"] += 1
            elif typ == "image":
                url = node["attrs"]["url"]
                if urlparse(url).scheme:
                    raise FormatError("Download and approve remote Markdown images before generation; use a local path")
                path = (self.source.parent / unquote(url)).resolve()
                if not path.is_file():
                    raise FormatError(f"Missing Markdown image: {path}")
                shape = p.add_run().add_picture(str(path))
                section = self.doc.sections[0]
                max_width = section.page_width - section.left_margin - section.right_margin
                if shape.width > max_width:
                    ratio = max_width / shape.width
                    shape.width = max_width
                    shape.height = round(shape.height * ratio)
                alt = self.plain(node.get("children", []))
                for prop in p._p.xpath(".//wp:docPr")[-1:]:
                    prop.set("descr", alt)
                self.counts["images"] += 1
            elif typ == "footnote_ref":
                self.ensure_notes()
                r = p.add_run()
                r.style = self.doc.styles["Footnote Reference"]
                r.font.superscript = True
                r._r.append(element("w:footnoteReference", id=node["attrs"]["index"]))
            elif typ == "inline_math":
                p._p.append(native_math(node["raw"]))
                self.counts["equations"] += 1
            else:
                raise FormatError(f"Unsupported inline element: {typ}")

    @staticmethod
    def plain(nodes):
        return "".join(n.get("raw", "") + MarkdownWriter.plain(n.get("children", [])) for n in nodes)

    def blocks(self, nodes, *, container=None, quote=False, list_id=None, depth=0, task_state=None):
        container = container or self.doc
        for node in nodes:
            typ = node["type"]
            if typ == "blank_line":
                continue
            if typ in {"paragraph", "block_text", "heading"}:
                if typ == "heading":
                    level = node["attrs"]["level"]
                    use_title = level == 1 and self.title_from_h1 and not self.title_used
                    style = "Title" if use_title else f"Heading {level}"
                    self.title_used = self.title_used or use_title
                    self.counts["headings"] += 1
                else:
                    style = "Quote" if quote else "Footnote Text" if isinstance(container, NoteContainer) else "Normal"
                p = container.add_paragraph(style=style)
                if typ == "heading" and node.get("_bookmark"):
                    p._p.append(element("w:bookmarkStart", id=self.bookmark_number, name=node["_bookmark"]))
                if task_state is not None:
                    w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
                    from lxml import etree
                    sdt = element("w:sdt")
                    pr = element("w:sdtPr")
                    check = etree.SubElement(pr, "{" + w14 + "}checkbox", nsmap={"w14": w14})
                    etree.SubElement(check, "{" + w14 + "}checked").set("{" + w14 + "}val", "1" if task_state else "0")
                    for tag, value in [("checkedState", "2612"), ("uncheckedState", "2610")]:
                        n = etree.SubElement(check, "{" + w14 + "}" + tag)
                        n.set("{" + w14 + "}val", value)
                        n.set("{" + w14 + "}font", "MS Gothic")
                    sdt.append(pr)
                    content, run, text = element("w:sdtContent"), element("w:r"), element("w:t")
                    text.text = "☒" if task_state else "☐"
                    run.append(text)
                    content.append(run)
                    sdt.append(content)
                    p._p.append(sdt)
                    p.add_run(" ")
                    p.paragraph_format.left_indent = Inches(0.25 * (depth + 1))
                    self.counts["task_boxes"] += 1
                self.inline(p, node.get("children", []))
                if typ == "heading" and node.get("_bookmark"):
                    p._p.append(element("w:bookmarkEnd", id=self.bookmark_number))
                    self.bookmark_number += 1
                if list_id is not None:
                    ppr = p._p.get_or_add_pPr()
                    numpr = element("w:numPr")
                    numpr.append(element("w:ilvl", val=depth))
                    numpr.append(element("w:numId", val=list_id))
                    ppr.append(numpr)
                    list_id = None  # subsequent blocks in one item are continuations
            elif typ == "list":
                if depth > 8:
                    raise FormatError("Word numbered lists support at most nine nesting levels")
                attrs = node.get("attrs", {})
                num_id = self.numbering(attrs.get("ordered", False), attrs.get("start", 1))
                for item in node["children"]:
                    first = True
                    task = item.get("attrs", {}).get("checked") if item["type"] == "task_list_item" else None
                    for block in item.get("children", []):
                        nested = block["type"] == "list"
                        self.blocks([block], container=container, quote=quote,
                                    list_id=num_id if first and not nested and task is None else None,
                                    depth=depth + 1 if nested else depth, task_state=task if first and not nested else None)
                        if not nested:
                            first = False
            elif typ == "block_quote":
                self.counts["quotes"] += 1
                self.blocks(node["children"], container=container, quote=True, depth=depth)
            elif typ == "block_code":
                p = container.add_paragraph(style="Code Block")
                p.add_run(node.get("raw", ""))
                self.counts["code_blocks"] += 1
            elif typ == "block_math":
                p = container.add_paragraph()
                p.paragraph_format.keep_together = True
                display = element("m:oMathPara")
                # m: attributes have a different namespace from w: properties.
                properties = element("m:oMathParaPr")
                alignment = element("m:jc")
                alignment.set(qn("m:val"), "center")
                properties.append(alignment)
                display.append(properties)
                display.append(native_math(node["raw"]))
                p._p.append(display)
                self.counts["equations"] += 1
            elif typ == "table":
                rows = []
                for block in node["children"]:
                    if block["type"] == "table_head":
                        rows.append(block["children"])
                    elif block["type"] == "table_body":
                        rows.extend(row["children"] for row in block["children"])
                if not rows or any(len(row) != len(rows[0]) for row in rows):
                    raise FormatError("Ragged Markdown table")
                table = container.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                table.rows[0]._tr.get_or_add_trPr().append(element("w:tblHeader"))
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        p = table.cell(i, j).paragraphs[0]
                        self.inline(p, cell.get("children", []))
                        align = cell.get("attrs", {}).get("align")
                        p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                                       "center": WD_ALIGN_PARAGRAPH.CENTER}.get(align)
                self.counts["tables"] += 1
            elif typ == "thematic_break":
                p = container.add_paragraph()
                border = element("w:pBdr")
                border.append(element("w:bottom", val="single", sz=4, color="808080"))
                p._p.get_or_add_pPr().append(border)
            elif typ == "footnotes":
                part = self.ensure_notes()
                for item in node["children"]:
                    note = element("w:footnote", id=item["attrs"]["index"])
                    part.element.append(note)
                    host = NoteContainer(note, part, self.doc)
                    self.blocks(item["children"], container=host)
                    first_p = note.find(qn("w:p"))
                    if first_p is None:
                        first_p = host.add_paragraph()._p
                        note.remove(first_p)
                        note.insert(0, first_p)
                    marker, ref = element("w:r"), element("w:footnoteRef")
                    mark_props = element("w:rPr")
                    mark_props.append(element("w:rStyle", val=self.doc.styles["Footnote Reference"].style_id))
                    mark_props.append(element("w:vertAlign", val="superscript"))
                    marker.append(mark_props)
                    marker.append(ref)
                    index = 1 if first_p.find(qn("w:pPr")) is not None else 0
                    first_p.insert(index, marker)
                    gap, gap_text = element("w:r"), element("w:t")
                    gap_text.text = " "
                    gap_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    gap.append(gap_text)
                    first_p.insert(index + 1, gap)
                    self.counts["footnotes"] += 1
            else:
                raise FormatError(f"Unsupported block: {typ}")


def generate_baseline(source, output, *, title_from_h1=False):
    if Path(source).resolve() == Path(output).resolve():
        raise FormatError("Do not overwrite Markdown input")
    parsed = parse_markdown(source)
    if parsed["unsupported"]:
        raise FormatError(f"Markdown requires a decision for unsupported elements: {parsed['unsupported']}")
    writer = MarkdownWriter(source, title_from_h1=title_from_h1)
    writer.prepare(parsed["ast"])
    writer.blocks(parsed["ast"])
    stream = BytesIO()
    writer.doc.save(stream)
    atomic_bytes(output, stream.getvalue())
    return {"source": parsed["source"], "source_sha256": parsed["sha256"], "output": str(Path(output).resolve()),
            "counts": writer.counts, "state": "semantic-baseline", "visual_qa": "not-verified",
            "note": "Apply a confirmed formatting plan to this baseline to produce tracked formatting revisions."}
