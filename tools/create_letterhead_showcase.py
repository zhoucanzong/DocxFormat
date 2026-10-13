"""Build original editable letterhead examples through the real component importer.

Outputs remain preview-only: this script never records a user's approval.
Run in a new directory; renderer and fonts are deliberately caller-controlled.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/docx-format/scripts"))
from docxformat.components import catalog, component_plan, preview_component
from docxformat.package import NS, Package, digest, element, qn, write_json

CASES = {
    "01-classic": ("常规红头", "青岚研究院文件"),
    "02-condensed": ("长名称与字距", "青岚城市发展研究中心文件"),
    "03-joint": ("联合发文组合", "青岚研究院／云川技术中心"),
    "04-outline": ("空心描边效果", "青岚研究院"),
}
RED = "C00000"


def font(run, size=16, *, red=False, bold=False, scale=100, spacing=0, outline=False):
    run.font.name = "Songti SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(RED if red else "000000")
    run.font.outline = outline
    rpr = run._r.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), "Songti SC")
    if scale != 100:
        rpr.append(element("w:w", val=str(scale)))
    if spacing:
        rpr.append(element("w:spacing", val=str(round(spacing * 20))))
    return run


def paragraph(doc, text="", *, size=16, align=WD_ALIGN_PARAGRAPH.LEFT,
              before=0, after=0, line=28, indent=False, bold=False, **kwargs):
    p = doc.add_paragraph()
    p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = Pt(line)
    if indent:
        fmt.first_line_indent = Pt(32)
    font(p.add_run(text), size, bold=bold, **kwargs)
    return p


def base():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(32), Mm(28)
    sec.left_margin, sec.right_margin = Mm(28), Mm(26)
    sec.footer_distance = Mm(15)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Songti SC", Pt(16)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Songti SC")
    title = doc.styles["Title"]
    title.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(title.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("原创虚构案例 · 仅作版式演示"), 9)
    return doc


def rule(p, *, double=False):
    borders = element("w:pBdr")
    borders.append(element("w:bottom", val="double" if double else "single",
                           sz="12", space="10", color=RED))
    p._p.get_or_add_pPr().append(borders)


def heading(doc, name, *, scale=100, size=36, spacing=0, outline=False):
    return paragraph(doc, name, size=size, red=True, bold=True, scale=scale,
                     spacing=spacing, outline=outline, line=54, after=10,
                     align=WD_ALIGN_PARAGRAPH.CENTER)


def joint_heading(doc):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Mm(110), Mm(46)
    props = table._tbl.tblPr
    borders = element("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element("w:" + side, val="nil"))
    props.append(borders)
    right = table.cell(0, 1).merge(table.cell(1, 1))
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for cell, text, size in [(table.cell(0, 0), "青岚研究院", 29),
                             (table.cell(1, 0), "云川技术中心", 29), (right, "文件", 38)]:
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = Pt(42)
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(text), size, red=True, bold=True, spacing=1)


def template(case):
    doc = base()
    if case == "03-joint":
        joint_heading(doc)
    elif case == "02-condensed":
        heading(doc, CASES[case][1], size=36, scale=78, spacing=1.2)
    elif case == "04-outline":
        heading(doc, CASES[case][1], size=42, spacing=4, outline=True)
        paragraph(doc, "QINGLAN  RESEARCH  INSTITUTE", size=10, red=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, line=16, after=8)
    else:
        heading(doc, CASES[case][1], spacing=1)
    p = paragraph(doc, "青研〔2026〕12号" if case != "03-joint" else "青研联〔2026〕3号",
                  size=13, align=WD_ALIGN_PARAGRAPH.CENTER, line=24, after=24)
    rule(p, double=case in {"03-joint", "04-outline"})
    # Deliberately unselected template body: migration must exclude this text.
    paragraph(doc, "模板占位正文，不得迁入。")
    return doc


def manuscript():
    doc = base()
    p = paragraph(doc, "关于开展文档规范化试点的通知", size=22,
                  align=WD_ALIGN_PARAGRAPH.CENTER, line=34, after=20)
    p.style = doc.styles["Title"]
    # Direct formatting above remains authoritative over theme-based Title style.
    paragraph(doc, "各项目组：", after=4)
    paragraph(doc, "为统一项目文档的编排方式，现开展文档规范化试点。试点先形成可审阅的样稿，经确认后再分批处理。有关安排如下。", indent=True)
    paragraph(doc, "一、先确认模板与适用范围", indent=True, bold=True, before=6)
    paragraph(doc, "各组提交现有文档及适用模板，标明需要保留的红头、目录、表格和特殊对象。遇到模板与文字要求不一致时，先核对再执行。", indent=True)
    paragraph(doc, "二、审阅后形成确认版本", indent=True, bold=True, before=6)
    paragraph(doc, "首份样稿通过后，按同组规则批量处理。格式修订保留待审状态；复杂红头迁入单独预览，原始文件另行保管。", indent=True)
    paragraph(doc, "青岚研究院", align=WD_ALIGN_PARAGRAPH.RIGHT, before=22)
    paragraph(doc, "2026年10月13日", align=WD_ALIGN_PARAGRAPH.RIGHT)
    return doc


def build(out):
    out = Path(out).resolve()
    if out.exists():
        raise SystemExit("Choose a new output directory; showcase files are not overwritten.")
    out.mkdir(parents=True)
    source = out / "manuscript.docx"
    manuscript().save(source)
    source_hash = digest(source.read_bytes())
    anchor = next(b for b in catalog(source)["blocks"] if b["part"] == "word/document.xml")
    results = []
    for case, (title, _) in CASES.items():
        folder = out / case
        folder.mkdir()
        path = folder / "template.docx"
        template(case).save(path)
        selected = [b for b in catalog(path)["blocks"] if b["part"] == "word/document.xml"][:-1]
        plan = component_plan(source, {"type": "letterhead", "template": str(path),
                              "block_ids": [b["id"] for b in selected], "before_id": anchor["id"],
                              "reason": "Original fictional documentation example; preview only"})
        write_json(folder / "plan.json", plan)
        preview = folder / "preview.docx"
        receipt = preview_component(plan, preview, folder / "receipt.json")
        # Independently assemble a visual reference without invoking the importer.
        reference = Package(source)
        body = reference.root("word/document.xml").find(qn("w:body"))
        src = Package(path)
        for i, block in enumerate(selected):
            body.insert(i, deepcopy(src.locate(block)))
        reference.touch("word/document.xml")
        reference.save(folder / "reference.docx")
        preview_pkg = Package(preview)
        text = "".join(preview_pkg.root("word/document.xml").xpath(".//w:t/text()", namespaces=NS))
        assert "模板占位正文" not in text
        assert digest(source.read_bytes()) == source_hash
        assert not receipt["verification"]["errors"]
        results.append({"case": case, "title": title, "preview": str(preview),
                        "state": receipt["state"], "source_unchanged": True,
                        "visual_qa": "not-verified"})
    write_json(out / "manifest.json", results)
    print(out)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=ROOT / "out/letterhead-showcase")
    build(parser.parse_args().out)
